#!/usr/bin/env python3
r"""
flag_missing_files_in_excel.py

Follow-up utility for the archive-extraction/organizing project. Takes the
error_log.csv produced by that pipeline and writes a marker into the
original Excel manifest for every row whose file was reported as
"not found", so you can see at a glance, right inside the sheet you
started from, which rows never got matched to a physical file.

------------------------------------------------------------------------
WHAT IT DOES
------------------------------------------------------------------------
1. Reads the error CSV. Auto-detects its 'filename' and 'reason' columns
   (delimiter is auto-detected too, since the pipeline's own logs use
   ';' while a hand-edited CSV might use ','). Filters down to just the
   rows whose 'reason' contains the phrase "file not found" (case-
   insensitive, so it matches wording like "file not found inside
   archive folder(s)" or "file not found inside unzipped folder(s)"
   equally). Rows about duplicates, ambiguous matches, or anything else
   are ignored -- they never contain that phrase, so they're excluded
   automatically without needing special-casing.
2. Builds a lookup of filename (case-insensitive, whitespace-trimmed) ->
   the exact reason text from the CSV. If the same filename appears in
   more than one "file not found" row with different reason text, the
   first one encountered is kept and a warning is logged so you can
   check the rest manually.
3. Opens the Excel file with openpyxl (keeps existing formatting/
   formulas intact) and locates the sheet's 'filename' column and its
   'file status' column, auto-detecting both by header text. If
   'file status' doesn't already exist, it's added as a new column at
   the end of the header row.
4. For every data row in the sheet whose filename matches an entry in
   the lookup, writes that row's reason text into 'file status'. Rows
   that don't match are left completely untouched (existing values, if
   any, are preserved).
5. Saves the result to a new file by default (so your original manifest
   is never touched) -- pass --in-place if you'd rather overwrite it.

------------------------------------------------------------------------
IMPORTANT NOTE ON MATCHING
------------------------------------------------------------------------
As requested, matching is done purely on filename. If your sheet has the
same filename appearing under more than one accessnumber/toegangsnummer,
ALL matching rows will be flagged, even though only one of them may
actually be the row the CSV error refers to -- the CSV/Excel pairing
doesn't currently carry enough information to disambiguate further. If
that turns out to be a problem in practice, matching on
(accessnumber, filename) together would be a straightforward addition --
just say so.

------------------------------------------------------------------------
USAGE
------------------------------------------------------------------------
    python flag_missing_files_in_excel.py \
        --csv "Q:\path\to\error_log.csv" \
        --excel "Q:\path\to\overview.xlsx"

Optional flags:
    --sheet NAME_OR_INDEX   which sheet to update (default: first sheet)
    --csv-filename-col      override: name of the filename column in the CSV
    --csv-reason-col        override: name of the reason column in the CSV
    --excel-filename-col    override: name of the filename column in Excel
    --status-col            name of the column to write into
                             (default: "file status")
    --output PATH           where to save the updated Excel file
                             (default: "<original>_updated.xlsx")
    --in-place              overwrite the original Excel file instead
"""

import argparse
import csv
import logging
import os
import re
import sys

import pandas as pd
from openpyxl import load_workbook

FILENAME_COL_CANDIDATES = ["filename", "file_name", "bestandsnaam", 'BESTANDSNAAM']
REASON_COL_CANDIDATES = ["reason"]
MATCH_PHRASE = "file not found"


def setup_logging():
    logger = logging.getLogger("flag_missing")
    logger.setLevel(logging.INFO)
    ch = logging.StreamHandler(sys.stdout)
    ch.setLevel(logging.INFO)
    ch.setFormatter(logging.Formatter("[%(levelname)s] %(message)s"))
    logger.addHandler(ch)
    return logger


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s).lower())


def find_column(columns, candidates):
    normalized = {col: norm(col) for col in columns}
    cand_norm = [norm(c) for c in candidates]
    for col, n in normalized.items():
        if n in cand_norm:
            return col
    return None


def normalize_filename(value):
    """Case-insensitive, whitespace-trimmed key used for matching."""
    return str(value).strip().lower()


# --------------------------------------------------------------------------
# step 1-2: read the CSV, build filename -> reason lookup
# --------------------------------------------------------------------------

def load_status_map(csv_path, filename_col_override, reason_col_override, logger):
    logger.info(f"Reading error CSV: {csv_path}")
    # sep=None + engine='python' auto-detects the delimiter (the pipeline's
    # own logs use ';', but a hand-edited CSV might use ',').
    df = pd.read_csv(csv_path, sep=None, engine="python", dtype=str, encoding="utf-8-sig")
    df.columns = [str(c).strip() for c in df.columns]

    filename_col = filename_col_override or find_column(df.columns, FILENAME_COL_CANDIDATES)
    reason_col = reason_col_override or find_column(df.columns, REASON_COL_CANDIDATES)

    missing = [name for name, col in [("filename", filename_col), ("reason", reason_col)] if col is None]
    if missing:
        raise SystemExit(
            f"Could not auto-detect column(s) in the CSV: {missing}. "
            f"Available columns are: {list(df.columns)}. "
            f"Re-run with --csv-filename-col / --csv-reason-col to specify them manually."
        )

    logger.info(f"Using CSV columns -> filename: '{filename_col}', reason: '{reason_col}'")

    df[filename_col] = df[filename_col].astype(str).str.strip()
    df[reason_col] = df[reason_col].astype(str).str.strip()

    mask = df[reason_col].str.contains(MATCH_PHRASE, case=False, na=False)
    matched = df[mask]
    logger.info(f"Found {len(matched)} row(s) in the CSV whose reason contains '{MATCH_PHRASE}' "
                f"(out of {len(df)} total error rows).")

    status_map = {}
    conflicts = 0
    for _, row in matched.iterrows():
        key = normalize_filename(row[filename_col])
        reason_text = row[reason_col]
        if key in status_map and status_map[key] != reason_text:
            conflicts += 1
            continue  # keep the first one encountered
        status_map.setdefault(key, reason_text)

    if conflicts:
        logger.warning(f"{conflicts} filename(s) had more than one differently-worded 'file not found' "
                        f"reason in the CSV; only the first occurrence was kept for each. Review manually "
                        f"if that matters for your case.")

    logger.info(f"Built a lookup covering {len(status_map)} distinct filename(s).")
    return status_map


# --------------------------------------------------------------------------
# step 3-5: write into the Excel file
# --------------------------------------------------------------------------

def resolve_sheet(wb, sheet_arg, logger):
    if sheet_arg is None:
        ws = wb.worksheets[0]
        logger.info(f"Using first sheet: '{ws.title}'")
        return ws
    # try as an integer index first, then fall back to a name lookup
    try:
        idx = int(sheet_arg)
        ws = wb.worksheets[idx]
        logger.info(f"Using sheet at index {idx}: '{ws.title}'")
        return ws
    except (ValueError, IndexError):
        pass
    if sheet_arg in wb.sheetnames:
        ws = wb[sheet_arg]
        logger.info(f"Using sheet: '{ws.title}'")
        return ws
    raise SystemExit(f"Sheet '{sheet_arg}' not found. Available sheets: {wb.sheetnames}")


def find_header_col(ws, header_row, candidates):
    """Returns the 1-based column index of the first header cell (in
    header_row) whose normalized text matches one of candidates, or
    None if not found."""
    cand_norm = [norm(c) for c in candidates]
    for cell in header_row:
        if cell.value is not None and norm(cell.value) in cand_norm:
            return cell.column
    return None


def update_excel(excel_path, sheet_arg, status_map, excel_filename_col_override,
                  status_col_name, output_path, logger):
    logger.info(f"Reading Excel file: {excel_path}")
    wb = load_workbook(excel_path)
    ws = resolve_sheet(wb, sheet_arg, logger)

    header_row = next(ws.iter_rows(min_row=1, max_row=1))

    if excel_filename_col_override:
        filename_col_idx = find_header_col(ws, header_row, [excel_filename_col_override])
    else:
        filename_col_idx = find_header_col(ws, header_row, FILENAME_COL_CANDIDATES)

    if filename_col_idx is None:
        headers = [c.value for c in header_row]
        raise SystemExit(
            f"Could not find a filename column in the Excel sheet's header row. "
            f"Available headers: {headers}. Re-run with --excel-filename-col to specify it manually."
        )

    status_col_idx = find_header_col(ws, header_row, [status_col_name])
    if status_col_idx is None:
        status_col_idx = ws.max_column + 1
        ws.cell(row=1, column=status_col_idx, value=status_col_name)
        logger.info(f"Added new column '{status_col_name}' at column {status_col_idx}.")
    else:
        logger.info(f"Column '{status_col_name}' already exists at column {status_col_idx}; reusing it.")

    matched_rows = 0
    for row in ws.iter_rows(min_row=2):
        filename_cell = row[filename_col_idx - 1]
        if filename_cell.value is None:
            continue
        key = normalize_filename(filename_cell.value)
        if key in status_map:
            ws.cell(row=filename_cell.row, column=status_col_idx, value=status_map[key])
            matched_rows += 1

    logger.info(f"Wrote '{status_col_name}' for {matched_rows} row(s) in the Excel sheet "
                f"(out of {len(status_map)} distinct filenames from the CSV -- some may not appear "
                f"in this sheet/tab at all, which is expected).")

    wb.save(output_path)
    logger.info(f"Saved: {output_path}")


# --------------------------------------------------------------------------
# main
# --------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(
        description="Flag 'file not found' rows from an error CSV into a 'file status' column of an Excel manifest."
    )
    parser.add_argument("--csv", required=True, help="Path to the error CSV")
    parser.add_argument("--excel", required=True, help="Path to the Excel manifest to update")
    parser.add_argument("--sheet", default=None, help="Sheet name or index to update (default: first sheet)")
    parser.add_argument("--csv-filename-col", default=None, help="Override: name of the filename column in the CSV")
    parser.add_argument("--csv-reason-col", default=None, help="Override: name of the reason column in the CSV")
    parser.add_argument("--excel-filename-col", default=None, help="Override: name of the filename column in Excel")
    parser.add_argument("--status-col", default="file status", help="Name of the column to write into (default: 'file status')")
    parser.add_argument("--output", default=None, help="Where to save the updated Excel file (default: '<original>_updated.xlsx')")
    parser.add_argument("--in-place", action="store_true", help="Overwrite the original Excel file instead of writing a new one")
    args = parser.parse_args()

    logger = setup_logging()

    if not os.path.isfile(args.csv):
        raise SystemExit(f"CSV file does not exist: {args.csv}")
    if not os.path.isfile(args.excel):
        raise SystemExit(f"Excel file does not exist: {args.excel}")

    if args.in_place:
        output_path = args.excel
    elif args.output:
        output_path = args.output
    else:
        stem, ext = os.path.splitext(args.excel)
        output_path = f"{stem}_updated{ext}"

    status_map = load_status_map(args.csv, args.csv_filename_col, args.csv_reason_col, logger)

    if not status_map:
        logger.warning(f"No rows matching '{MATCH_PHRASE}' were found in the CSV -- nothing to write. "
                        f"Double-check --csv-reason-col if this seems unexpected.")

    update_excel(args.excel, args.sheet, status_map, args.excel_filename_col,
                 args.status_col, output_path, logger)


if __name__ == "__main__":
    main()
